from io import BytesIO

from openpyxl import load_workbook

from hrc_reconcile.models import ComparisonCell, ComparisonRow, FieldStatus, ReviewSession, ToleranceConfig
from hrc_reconcile.reporting import build_report_workbook


def test_report_has_expected_sheets_values_and_status_colors():
    rows = [ComparisonRow("355818|3180730000", (
        ComparisonCell("Heat no", "355818", "355818", "355818", "355818", None, 0, FieldStatus.MATCH, 0.99),
        ComparisonCell("C", 7, 8, 7, 8, 1, 0.01, FieldStatus.MISMATCH, 0.99, "outside tolerance"),
    ), FieldStatus.MISMATCH)]
    workbook = build_report_workbook(rows, ReviewSession(confirmed=True), ToleranceConfig({"C": 0.01}))
    book = load_workbook(BytesIO(workbook.getvalue()))

    assert {"Summary", "Row results", "Cell details", "Confirmed OCR", "Tolerances"}.issubset(book.sheetnames)
    assert book["Cell details"]["H3"].value == "mismatch"
    assert book["Cell details"]["H3"].fill.fill_type == "solid"
    assert book["Cell details"]["H3"].fill.fgColor.rgb.endswith("F4CCCC")
