from __future__ import annotations

from collections import Counter
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from .models import ComparisonRow, FieldStatus, ReviewSession, ToleranceConfig


_COLORS = {
    FieldStatus.MATCH: "C6EFCE",
    FieldStatus.MISMATCH: "F4CCCC",
    FieldStatus.MISSING_ON_DOCUMENT: "FFF2CC",
    FieldStatus.MISSING_IN_COMPANY_DATA: "D9EAF7",
    FieldStatus.OCR_UNREADABLE: "FCE4D6",
    FieldStatus.MANUAL_REVIEW: "E4DFEC",
}


def _style_sheet(sheet) -> None:
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F4E78")
    for column in range(1, sheet.max_column + 1):
        sheet.column_dimensions[get_column_letter(column)].width = min(max(len(str(sheet.cell(1, column).value or "")) + 4, 12), 32)


def build_report_workbook(rows: list[ComparisonRow], review_session: ReviewSession, tolerances: ToleranceConfig) -> BytesIO:
    workbook = Workbook()
    summary = workbook.active
    summary.title = "Summary"
    statuses = Counter(row.status.value for row in rows)
    summary.append(["Status", "Rows"])
    for status, count in sorted(statuses.items()):
        summary.append([status, count])

    row_sheet = workbook.create_sheet("Row results")
    row_sheet.append(["Row key", "Status", "Note"])
    for row in rows:
        row_sheet.append([row.row_key, row.status.value, row.note])
        row_sheet.cell(row_sheet.max_row, 2).fill = PatternFill("solid", fgColor=_COLORS[row.status])

    details = workbook.create_sheet("Cell details")
    details.append(["Row key", "Field", "Company value", "Document value", "Normalized company", "Normalized document", "Difference", "Status", "Tolerance", "OCR confidence", "Note"])
    for row in rows:
        for cell in row.cells:
            details.append([row.row_key, cell.field, cell.company_value, cell.document_value, cell.normalized_company_value, cell.normalized_document_value, cell.difference, cell.status.value, cell.tolerance, cell.ocr_confidence, cell.note])
            details.cell(details.max_row, 8).fill = PatternFill("solid", fgColor=_COLORS[cell.status])

    ocr_sheet = workbook.create_sheet("Confirmed OCR")
    ocr_sheet.append(["Row key", "Field", "Raw text", "Value", "Confidence", "Manually edited"])
    for cell in review_session.cells:
        ocr_sheet.append([cell.row_key, cell.field, cell.raw_text, cell.value, cell.confidence, cell.manually_edited])

    tolerance_sheet = workbook.create_sheet("Tolerances")
    tolerance_sheet.append(["Field", "Tolerance"])
    for field, tolerance in sorted(tolerances.values.items()):
        tolerance_sheet.append([field, tolerance])

    for sheet in workbook.worksheets:
        _style_sheet(sheet)
    output = BytesIO()
    workbook.save(output)
    output.seek(0)
    return output
